"""Insert the mechanism summary table after PowerPoint has closed."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from .config import BODY_LEFT, BODY_TOP, BODY_WIDTH, MECHANISM, SLIDE_H

INK = RGBColor(0x1F, 0x29, 0x37)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _pt(value: float) -> int:
    return int(Emu(value * 12700))


def _fill(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for child in list(tc_pr):
        if child.tag == qn("a:solidFill"):
            tc_pr.remove(child)
    tc_pr.append(
        parse_xml(
            '<a:solidFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            f'<a:srgbClr val="{hex_color}"/></a:solidFill>'
        )
    )


def _write(cell, text: str, *, size: float, bold: bool, color) -> None:
    frame = cell.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Pt(3)
    frame.margin_right = Pt(3)
    frame.margin_top = Pt(2)
    frame.margin_bottom = Pt(2)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Malgun Gothic"
    run.font.color.rgb = color


def _slide_title(slide) -> str:
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            return shape.text_frame.text
    return ""


def _summary_rows() -> list[tuple[str, ...]]:
    rows: list[tuple[str, ...]] = [
        ("Arm", "Cell", "mechanism", "Q_cliff_abs", "note"),
    ]
    notes = {
        ("SJ900_dry", "M01Ch022"): "Si only fade closest (H1)",
        ("SJ900_dry", "M01Ch024"): "Gr interval also shrinks",
        ("SJ900_dry", "M01Ch025"): "short life, tagged ~134",
        ("SJ1300_dry", "M01Ch010"): "cliff shorter than SJ900",
        ("SJ1300_dry", "M01Ch011"): "",
        ("SJ1300_dry", "M01Ch012"): "cliff null",
    }
    if not MECHANISM.exists():
        rows.append(("—", "—", "file missing", "—", str(MECHANISM)))
        return rows
    df = pd.read_csv(MECHANISM)
    for _, row in df.iterrows():
        arm = str(row.get("arm", ""))
        cell = str(row.get("cell_id", ""))
        mech = str(row.get("mechanism_state", ""))
        cliff = row.get("Q_cliff_abs")
        cliff_s = f"{float(cliff):.1f} Ah" if pd.notna(cliff) else "—"
        rows.append((arm.replace("_dry", ""), cell.replace("M01", ""), mech, cliff_s, notes.get((arm, cell), "")))
    return rows


def add_summary_table(path: Path) -> None:
    prs = Presentation(str(path))
    target = None
    for slide in prs.slides:
        if "판정" in _slide_title(slide) or "summary" in _slide_title(slide).casefold():
            target = slide
            break
    if target is None:
        target = prs.slides[-1]
    rows = _summary_rows()
    left = _pt(BODY_LEFT)
    top = _pt(BODY_TOP)
    width = _pt(BODY_WIDTH)
    height = _pt(min(220.0, (SLIDE_H - BODY_TOP - 40)))
    table = target.shapes.add_table(len(rows), 5, left, top, width, height).table
    widths = [0.14, 0.14, 0.22, 0.16, 0.34]
    for i, frac in enumerate(widths):
        table.columns[i].width = _pt(BODY_WIDTH * frac)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            if r == 0:
                _write(cell, val, size=11, bold=True, color=WHITE)
                _fill(cell, "15325B")
            else:
                _write(cell, val, size=11, bold=False, color=INK)
                _fill(cell, "F4F7FA" if r % 2 else "FFFFFF")
    tmp = path.with_suffix(".tmp.pptx")
    prs.save(tmp)
    tmp.replace(path)


def add_last_value_table(path: Path, tagged: dict, metrics) -> None:
    """Last tagged-cycle value per cell for the selected overlay metrics."""
    from .config import OVERLAY_CELLS

    headers = ["Arm", "Cell", "t#"] + [m.column for m in metrics]
    rows: list[tuple[str, ...]] = [tuple(headers)]
    for arm, cells in OVERLAY_CELLS.items():
        df = tagged.get(arm)
        if df is None:
            continue
        for cell in cells:
            g = df[df["cell_id"] == cell]
            if g.empty:
                continue
            last = g.sort_values("tagged_cycle").iloc[-1]
            tidx = last.get("tagged_cycle", "")
            values = [arm, str(cell).replace("M01", ""), str(int(tidx) if pd.notna(tidx) else "")]
            for item in metrics:
                val = last[item.column] if item.column in last.index else None
                if val is None or pd.isna(val):
                    values.append("—")
                else:
                    number = float(val)
                    values.append(f"{number:.4g}" if abs(number) < 1000 else f"{number:.1f}")
            rows.append(tuple(values))

    prs = Presentation(str(path))
    target = None
    for slide in prs.slides:
        title = _slide_title(slide)
        if "말기" in title or "last" in title.casefold() or "요약" in title:
            target = slide
            break
    if target is None:
        target = prs.slides[-1]
    n_col = len(headers)
    left = _pt(BODY_LEFT)
    top = _pt(BODY_TOP)
    width = _pt(BODY_WIDTH)
    height = _pt(min(280.0, (SLIDE_H - BODY_TOP - 30)))
    table = target.shapes.add_table(len(rows), n_col, left, top, width, height).table
    first_w = 0.12 + 0.12 + 0.08
    rest = (1.0 - first_w) / max(n_col - 3, 1)
    widths = [0.12, 0.12, 0.08] + [rest] * (n_col - 3)
    for i, frac in enumerate(widths):
        table.columns[i].width = _pt(BODY_WIDTH * frac)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            if r == 0:
                _write(cell, val, size=9, bold=True, color=WHITE)
                _fill(cell, "15325B")
            else:
                _write(cell, val, size=10, bold=False, color=INK)
                _fill(cell, "F4F7FA" if r % 2 else "FFFFFF")
    tmp = path.with_suffix(".tmp.pptx")
    prs.save(tmp)
    tmp.replace(path)
