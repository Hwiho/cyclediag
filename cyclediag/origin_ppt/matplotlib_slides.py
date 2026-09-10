"""Origin-free matplotlib PowerPoint for overlay decks."""

from __future__ import annotations

from pathlib import Path
import tempfile

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from .config import SLIDE_H, SLIDE_W
from .tables import add_last_value_table, add_summary_table

INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x4B, 0x55, 0x63)


def _pt_to_emu(value: float) -> int:
    return int(Emu(value * 12700))


def _font(run, *, size: float, bold: bool = False, color=INK) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Malgun Gothic"


def _header(slide, title: str, subtitle: str) -> None:
    box = slide.shapes.add_textbox(_pt_to_emu(20), _pt_to_emu(8), _pt_to_emu(920), _pt_to_emu(28))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = title
    _font(run, size=18, bold=True)
    note = slide.shapes.add_textbox(_pt_to_emu(20), _pt_to_emu(34), _pt_to_emu(920), _pt_to_emu(16))
    p = note.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = subtitle
    _font(run, size=11, bold=False, color=MUTED)


def _plot_graphs(spec: dict, png: Path) -> None:
    graphs = spec["graphs"]
    layout = spec["layout"]
    if layout == "quad":
        rows, cols = 2, 2
    elif layout == "triple":
        rows, cols = 1, 3
    else:
        rows, cols = 1, max(len(graphs), 1)
        cols = len(graphs)
    fig, axes = plt.subplots(rows, cols, figsize=(13.2, 5.6), squeeze=False)
    flat = list(axes.ravel())
    for ax in flat[len(graphs) :]:
        ax.axis("off")
    for ax, graph in zip(flat, graphs):
        for series in graph["series"]:
            style = "o" if series.get("kind") == "s" else "-"
            ax.plot(
                series["x"],
                series["y"],
                style,
                color=series.get("color", "#333333"),
                label=series["label"],
                linewidth=1.6,
                markersize=4,
            )
        ax.set_xlabel(graph["x_title"])
        ax.set_ylabel(graph["y_title"])
        ax.set_xlim(*graph["x_limits"])
        ax.set_ylim(*graph["y_limits"])
        ax.legend(fontsize=7, frameon=False, loc="best")
        ax.grid(True, alpha=0.25)
    fig.tight_layout()
    fig.savefig(png, dpi=140)
    plt.close(fig)


def build_matplotlib(slides: list[dict], out_ppt: Path, *, deck, tagged: dict) -> None:
    out_ppt.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W / 72.0)
    prs.slide_height = Inches(SLIDE_H / 72.0)
    blank = prs.slide_layouts[6]
    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        for index, spec in enumerate(slides, start=1):
            slide = prs.slides.add_slide(blank)
            _header(slide, spec["title"], spec["subtitle"])
            if spec["kind"] in {"cover", "text"}:
                body = spec.get("body") or spec["subtitle"]
                box = slide.shapes.add_textbox(
                    _pt_to_emu(38), _pt_to_emu(92), _pt_to_emu(884), _pt_to_emu(380)
                )
                tf = box.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.LEFT
                run = p.add_run()
                run.text = body.replace("\r", "\n")
                _font(run, size=16)
                continue
            if spec["kind"] != "graphs":
                continue
            png = tmp_dir / f"slide_{index:02d}.png"
            _plot_graphs(spec, png)
            slide.shapes.add_picture(
                str(png),
                _pt_to_emu(20),
                _pt_to_emu(88),
                width=_pt_to_emu(920),
                height=_pt_to_emu(410),
            )
        prs.save(str(out_ppt))
    if deck.summary == "mechanism":
        add_summary_table(out_ppt)
    elif deck.summary == "last_values":
        add_last_value_table(out_ppt, tagged, deck.metrics)
    print(f"Saved matplotlib slides={len(slides)} -> {out_ppt}", flush=True)
